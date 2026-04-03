import sys
from docx import Document

file_path = r'c:\Users\ASUS\Projects\CuraNex\Problem Statement.docx'
doc = Document(file_path)

# 1. Update Technical challenges
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("Shocks — "):
        p1 = p.insert_paragraph_before("Recursive Auto-Regressive Smoothing — Multi-step forecasting algorithms naturally lose variance and decay into averages over extended horizons. Preserving granular historical spikes across long-term sequences requires mathematically fusing AI trend components with static statistical memories (e.g., Seasonal Naïve baselines).")
        p1.style = p.style
        
        p2 = p.insert_paragraph_before("Large-Scale Memory Exhaustion — Recursively shifting and calculating deep temporal lag matrices across hundreds of thousands of panel rows exhausts operational memory (RAM), requiring sophisticated sliding-window memory slicing and rigorous garbage collection architectures to execute continuous inference sequences safely.")
        p2.style = p.style
        break

# 2. Update Forecasting Integration
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("API — Allows other systems"):
        p3 = p.insert_paragraph_before("Future Extrapolation Intelligence — The pipeline intrinsically bridges mathematical horizons (out to 17+ weeks / 4+ months) by securely generating isolated \"pure future\" inference matrices initialized with projected holiday calendars and exogenous health vectors. The Streamlit dashboard visually stitches and renders these deep-future extrapolated trajectories atop the historical anchors so planners can evaluate long-term supply pipelines continuously.")
        p3.style = p.style
        break

# 3. Fix Inaccuracies (Updating the wording to reflect the completed recursive generation)
for i, p in enumerate(doc.paragraphs):
    if "Extending to explicit multi-step targets per horizon is a natural production enhancement" in p.text:
        p.text = p.text.replace(
            "Extending to explicit multi-step targets per horizon is a natural production enhancement without changing the overall methodology.", 
            "The pipeline now natively executes recursive autoregressive feedback loops on synthetic future structures to dynamically extrapolate continuous multi-step targets over multiple consecutive months."
        )

doc.save(file_path)
print("Document updated successfully.")
